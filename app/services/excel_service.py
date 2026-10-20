"""Bounded, read-only XLSX parsing without executing formulas or retaining uploads."""

from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.models.schemas import RecipientRequest
from app.services.text_layout_service import RecipientText

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_DATA_ROWS = 100
HEADERS = ("Name", "Second_Line", "City")


class ExcelError(ValueError):
    """A safe workbook-level error suitable for a client response."""


@dataclass(frozen=True)
class ExcelRow:
    number: int
    recipient: RecipientText | None = None
    error: str | None = None
    skipped: bool = False
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class WorkbookRows:
    rows: tuple[ExcelRow, ...]

    @property
    def total_rows(self) -> int:
        return len(self.rows)


def parse_workbook(data: bytes, filename: str) -> WorkbookRows:
    if not filename.lower().endswith(".xlsx"):
        raise ExcelError("Upload an Excel .xlsx file.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExcelError("Excel uploads must not exceed 5 MB.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 256 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ExcelError("Workbook contents exceed the supported size.")
            if any(entry.flag_bits & 1 or entry.filename.lower().endswith("vbaproject.bin") for entry in entries):
                raise ExcelError("Encrypted or macro-enabled workbooks are not supported.")
    except (BadZipFile, OSError, ValueError) as error:
        if isinstance(error, ExcelError):
            raise
        raise ExcelError("The file is not a readable .xlsx workbook.") from None

    workbook = None
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if len(workbook.sheetnames) != 1:
            raise ExcelError("Use a workbook with exactly one worksheet.")
        worksheet = workbook.worksheets[0]
        worksheet.reset_dimensions()
        iterator = worksheet.iter_rows()
        header = next(iterator, ())
        if tuple(cell.value for cell in header) != HEADERS:
            raise ExcelError("The first row must contain exactly: Name, Second_Line, City.")
        rows = []
        for number, cells in enumerate(iterator, start=2):
            if number > MAX_DATA_ROWS + 1:
                raise ExcelError("A batch may contain at most 100 data rows, including blank rows.")
            if any(cell.value is not None for cell in cells[3:]):
                raise ExcelError("Only the Name, Second_Line, and City columns are supported.")
            padded = list(cells[:3])
            name = padded[0].value if padded else None
            if name is None or isinstance(name, str) and not name.strip():
                rows.append(ExcelRow(number, skipped=True))
                continue
            if any(cell.data_type in ("f", "e") for cell in padded):
                rows.append(ExcelRow(number, error="Formula or Excel error cells are not supported; replace them with text."))
                continue
            values = [cell.value if cell.value is not None else "" for cell in padded]
            values += [""] * (3 - len(values))
            try:
                recipient = RecipientRequest(name=values[0], second_line=values[1], city=values[2])
            except ValidationError:
                rows.append(ExcelRow(number, error="Use text cells with at most 500 characters; Name is required."))
                continue
            rows.append(ExcelRow(number, recipient=RecipientText(
                recipient.name, recipient.second_line, recipient.city
            )))
        return WorkbookRows(tuple(rows))
    except ExcelError:
        raise
    except Exception:
        raise ExcelError("The file is not a readable .xlsx workbook.") from None
    finally:
        if workbook is not None:
            workbook.close()


def example_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recipients"
    sheet.append(HEADERS)
    sheet.append(("अमित जैन", "सपरिवार", "जयपुर"))
    sheet.append(("Amit Jain", "", "Jaipur"))
    sheet.append(("Neha Jain", "", ""))
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 25
    sheet.column_dimensions["C"].width = 25
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()